"""
MCP Filesystem Server — workspace tools for AI agents
─────────────────────────────────────────────────────
Layer A: plain Python functions  → imported directly by app.py (no async overhead)
Layer B: FastMCP decorators      → run as standalone MCP server via: python mcp_server.py
"""
import os
from pathlib import Path
from datetime import datetime

# ─── Layer A: Core filesystem functions ───────────────────────────────────────


def _validate_path(workspace: str, filename: str) -> str:
    """Resolve and validate that filename stays within workspace.
    Raises ValueError on path traversal attempt. Returns safe absolute path."""
    workspace_abs = str(Path(workspace).resolve())
    target = str(Path(workspace_abs, filename).resolve())
    try:
        common = os.path.commonpath([workspace_abs, target])
        if common != workspace_abs:
            raise ValueError(f"ไม่อนุญาต: '{filename}' อยู่นอก workspace")
    except ValueError:
        raise ValueError(f"ไม่อนุญาต: '{filename}' อยู่นอก workspace")
    return target


def fs_list_files(workspace: str) -> list:
    """List all files (non-recursive) in workspace.
    Returns list of {name, size, modified} dicts."""
    workspace_abs = os.path.abspath(workspace)
    if not os.path.isdir(workspace_abs):
        return []
    result = []
    try:
        for entry in sorted(os.scandir(workspace_abs), key=lambda e: e.name):
            if entry.is_file():
                stat = entry.stat()
                result.append({
                    'name': entry.name,
                    'size': stat.st_size,
                    'modified': datetime.fromtimestamp(stat.st_mtime).strftime('%Y-%m-%d %H:%M')
                })
    except OSError:
        pass
    return result


def fs_create_file(workspace: str, filename: str, content: str) -> str:
    """Create a new file in workspace. Raises FileExistsError if already exists."""
    path = _validate_path(workspace, filename)
    if os.path.exists(path):
        raise FileExistsError(f"ไฟล์ '{filename}' มีอยู่แล้ว ใช้ update_file เพื่อแก้ไข")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(content)
    return f"✅ สร้างไฟล์ '{filename}' สำเร็จ ({len(content)} ตัวอักษร)"


def fs_read_file(workspace: str, filename: str) -> str:
    """Read file contents from workspace. Binary formats (docx, xlsx, pdf) are extracted to plain text."""
    path = _validate_path(workspace, filename)
    if not os.path.exists(path):
        raise FileNotFoundError(f"ไม่พบไฟล์ '{filename}'")
    ext = Path(path).suffix.lower()
    if ext == '.docx':
        try:
            from docx import Document
            doc = Document(path)
            text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
            return text or '(ไฟล์ว่างเปล่า)'
        except Exception as e:
            raise ValueError(f"ไม่สามารถอ่านไฟล์ .docx ได้: {e}")
    if ext == '.xlsx':
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            lines = []
            for sheet in wb.worksheets:
                lines.append(f"[Sheet: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    line = '\t'.join('' if v is None else str(v) for v in row)
                    if line.strip():
                        lines.append(line)
            wb.close()
            return '\n'.join(lines) or '(ไฟล์ว่างเปล่า)'
        except Exception as e:
            raise ValueError(f"ไม่สามารถอ่านไฟล์ .xlsx ได้: {e}")
    if ext == '.pdf':
        try:
            import pdfplumber
            lines = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        lines.append(t)
            return '\n\n'.join(lines) or '(ไฟล์ว่างเปล่า)'
        except Exception as e:
            raise ValueError(f"ไม่สามารถอ่านไฟล์ .pdf ได้: {e}")
    _MAX_CHARS = 80_000  # ~20K tokens — safe for most LLM context windows
    with open(path, 'r', encoding='utf-8') as f:
        content = f.read(_MAX_CHARS + 1)
    if len(content) > _MAX_CHARS:
        content = content[:_MAX_CHARS]
        content += f'\n\n[⚠️ ไฟล์ถูกตัดที่ {_MAX_CHARS:,} ตัวอักษร เนื่องจากไฟล์มีขนาดใหญ่เกินไป]'
    return content


def fs_update_file(workspace: str, filename: str, content: str) -> str:
    """Overwrite existing file in workspace with new content."""
    path = _validate_path(workspace, filename)
    if not os.path.exists(path):
        raise FileNotFoundError(f"ไม่พบไฟล์ '{filename}' กรุณาใช้ create_file สำหรับสร้างไฟล์ใหม่")
    with open(path, 'w', encoding='utf-8') as f:
        f.write(content)
    return f"✅ อัปเดตไฟล์ '{filename}' สำเร็จ ({len(content)} ตัวอักษร)"


def fs_delete_file(workspace: str, filename: str) -> str:
    """Delete a file from workspace."""
    path = _validate_path(workspace, filename)
    if not os.path.exists(path):
        raise FileNotFoundError(f"ไม่พบไฟล์ '{filename}'")
    os.remove(path)
    return f"✅ ลบไฟล์ '{filename}' สำเร็จ"


# ─── Layer B: FastMCP server (run standalone) ─────────────────────────────────

try:
    from mcp.server.fastmcp import FastMCP

    _WORKSPACE = os.getenv(
        "WORKSPACE_PATH",
        os.path.join(os.path.dirname(os.path.abspath(__file__)), "workspace")
    )

    mcp = FastMCP("workspace-fs")

    @mcp.tool()
    def create_file(filename: str, content: str) -> str:
        """สร้างไฟล์ใหม่ใน workspace พร้อมเนื้อหา"""
        return fs_create_file(_WORKSPACE, filename, content)

    @mcp.tool()
    def read_file(filename: str) -> str:
        """อ่านเนื้อหาของไฟล์จาก workspace"""
        return fs_read_file(_WORKSPACE, filename)

    @mcp.tool()
    def update_file(filename: str, content: str) -> str:
        """แก้ไขเนื้อหาของไฟล์ที่มีอยู่แล้วใน workspace"""
        return fs_update_file(_WORKSPACE, filename, content)

    @mcp.tool()
    def delete_file(filename: str) -> str:
        """ลบไฟล์จาก workspace"""
        return fs_delete_file(_WORKSPACE, filename)

    @mcp.tool()
    def list_files() -> str:
        """แสดงรายการไฟล์ทั้งหมดที่มีอยู่ใน workspace"""
        files = fs_list_files(_WORKSPACE)
        if not files:
            return "workspace ว่างเปล่า ยังไม่มีไฟล์"
        return "\n".join(
            f"- {f['name']} ({f['size']} bytes, แก้ไขล่าสุด: {f['modified']})"
            for f in files
        )

    if __name__ == "__main__":
        print(f"🚀 MCP Filesystem Server เริ่มต้นแล้ว")
        print(f"📁 Workspace: {_WORKSPACE}")
        mcp.run()

except ImportError:
    if __name__ == "__main__":
        print("❌ กรุณาติดตั้ง mcp package: pip install mcp")
